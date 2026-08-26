"""Midday attendance digest — one email per company, every day.

Ingest that fails quietly is worse than ingest that fails loudly: a bridge
that stopped three days ago looks exactly like a workforce that stopped
punching. This job puts the day's data-quality problems in front of a human
while they are still fixable, and proves the devices are alive when they are.

Three sections, each answering a different question:

  A. Missed Punches (yesterday)
     Employees whose punch count for the day is ODD. An odd count means a
     punch is missing — someone forgot to tap out, or tapped in twice. The
     day still produces an Attendance row (desired_log_type forces the last
     punch to OUT precisely so it does), but the hours are wrong, so these
     are the rows worth correcting in Attendance Fix before payroll.
     Yesterday, not today: at send time no day shift has ended, so today's
     counts are all legitimately odd and would bury the real ones.

  B. Not Seen Today (as of send time)
     Employees rostered on a shift that has already started, with no punch
     today, no approved leave, and not on holiday. These are the people to
     chase now — while the day can still be corrected — rather than the
     morning payroll runs.

  C. Device Status (live)
     Last contact per device against its own alert threshold. The heartbeat
     job already emails on the down/recovered EDGE; this restates the state
     daily, so a device that went down and was acknowledged but never fixed
     cannot fade out of view.

The digest follows the working calendar, not the wall calendar. Nothing is
sent on a company holiday — nobody is expected to punch, so every section
would either be empty or wrong — and section A reports the last WORKING day
rather than literally yesterday. Reporting "yesterday" would silently lose
Friday: Saturday is a weekly off, so Sunday's digest would look at Saturday
and find nothing, and Friday's missed punches would never reach anyone.

Recipients come from each device's `digest_recipients`, pooled per company —
one email per company, not per device. A company with no recipients
configured is skipped silently; that is how you turn the digest off.
"""

from collections import defaultdict
from datetime import datetime, time, timedelta

import frappe
from frappe import _
from frappe.utils import add_days, cint, flt, get_datetime, getdate, now_datetime

# Inline tables are trimmed to keep the email readable on a phone; the full
# lists always travel in the attached workbook.
MAX_ROWS_INLINE = 15


# ---------------------------------------------------------------------------
# scheduler entry point
# ---------------------------------------------------------------------------

def send_daily_attendance_digest():
	"""Cron entry point (see hooks.scheduler_events). One email per company."""
	companies = frappe.get_all(
		"Biometric Device",
		filters={"enabled": 1},
		fields=["company"],
		group_by="company",
		pluck="company",
	)
	if not companies:
		return

	log = frappe.logger("biometric")
	for company in companies:
		if not company:
			continue
		try:
			recipients = get_digest_recipients(company)
			if not recipients:
				continue
			holiday_list = get_company_holiday_list(company)
			if _is_holiday(holiday_list, getdate(now_datetime())):
				log.info("attendance digest skipped: company=%s is on holiday today", company)
				continue
			digest = build_digest(company, holiday_list=holiday_list)
			send_digest_email(company, recipients, digest)
			log.info(
				"attendance digest sent: company=%s to=%s missed=%d unseen=%d devices=%d",
				company,
				",".join(recipients),
				len(digest["missed_punches"]),
				len(digest["not_seen"]),
				len(digest["devices"]),
			)
		except Exception:
			frappe.log_error(
				title=f"Attendance digest failed: {company}",
				message=frappe.get_traceback(),
			)


def get_digest_recipients(company):
	"""Pooled, de-duplicated recipient emails across the company's devices."""
	rows = frappe.db.sql(
		"""
		SELECT DISTINCT r.email
		FROM `tabBiometric Device Alert Recipient` r
		INNER JOIN `tabBiometric Device` d ON d.name = r.parent
		WHERE r.parenttype = 'Biometric Device'
		  AND r.parentfield = 'digest_recipients'
		  AND d.company = %s
		  AND d.enabled = 1
		""",
		company,
		as_dict=True,
	)
	return [r.email for r in rows if r.email]


# ---------------------------------------------------------------------------
# data
# ---------------------------------------------------------------------------

def build_digest(company, for_date=None, as_of=None, holiday_list=None):
	"""Assemble the three sections.

	`for_date` is the day section A reports on — by default the last WORKING
	day, not literally yesterday, so a Sunday digest reports Friday instead of
	an empty Saturday. `as_of` is the moment section B judges (default now)."""
	as_of = get_datetime(as_of) if as_of else now_datetime()
	today = getdate(as_of)
	if holiday_list is None:
		holiday_list = get_company_holiday_list(company)
	for_date = getdate(for_date) if for_date else last_working_day(holiday_list, today)

	return {
		"company": company,
		"for_date": for_date,
		"today": today,
		"as_of": as_of,
		"holiday_list": holiday_list,
		"missed_punches": _missed_punches(company, for_date),
		"not_seen": _not_seen_today(company, today, as_of, holiday_list),
		"devices": _device_status(company, as_of),
	}


def _missed_punches(company, for_date):
	"""Employees with an ODD number of punches on `for_date`."""
	day_start = datetime.combine(for_date, time.min)
	day_end = datetime.combine(for_date, time.max)

	rows = frappe.db.sql(
		"""
		SELECT
			c.employee,
			e.employee_name,
			e.department,
			COUNT(*)            AS punches,
			MIN(c.`time`)       AS first_punch,
			MAX(c.`time`)       AS last_punch
		FROM `tabEmployee Checkin` c
		INNER JOIN `tabEmployee` e ON e.name = c.employee
		WHERE e.company = %(company)s
		  AND c.`time` BETWEEN %(start)s AND %(end)s
		GROUP BY c.employee, e.employee_name, e.department
		HAVING COUNT(*) %% 2 = 1
		ORDER BY COUNT(*) ASC, e.employee_name ASC
		""",
		{"company": company, "start": day_start, "end": day_end},
		as_dict=True,
	)

	for r in rows:
		att = frappe.db.get_value(
			"Attendance",
			{"employee": r.employee, "attendance_date": for_date, "docstatus": ("<", 2)},
			["name", "status", "working_hours", "shift"],
			as_dict=True,
		)
		r.attendance = att.name if att else None
		r.status = att.status if att else "— none —"
		r.working_hours = flt(att.working_hours, 2) if att else 0
		r.shift = (att.shift if att else None) or ""
	return rows


def _not_seen_today(company, today, as_of, company_holiday_list=None):
	"""Rostered, shift already started, no punch, no approved leave, not a holiday.

	The holiday list is resolved shift -> employee -> company, matching what
	HRMS auto-attendance honours. Employee.holiday_list alone is not enough:
	on this site 1 of 299 employees has one, so trusting it would report the
	whole roster as missing every Saturday."""
	assignments = frappe.db.sql(
		"""
		SELECT sa.employee, e.employee_name, e.department,
		       COALESCE(NULLIF(st.holiday_list, ''), NULLIF(e.holiday_list, '')) AS holiday_list,
		       sa.shift_type, st.start_time
		FROM `tabShift Assignment` sa
		INNER JOIN `tabEmployee`   e  ON e.name = sa.employee
		INNER JOIN `tabShift Type` st ON st.name = sa.shift_type
		WHERE sa.docstatus = 1
		  AND sa.status = 'Active'
		  AND e.company = %(company)s
		  AND e.status = 'Active'
		  AND sa.start_date <= %(today)s
		  AND (sa.end_date IS NULL OR sa.end_date = '' OR sa.end_date >= %(today)s)
		""",
		{"company": company, "today": today},
		as_dict=True,
	)
	if not assignments:
		return []

	# Only judge someone once their shift has actually begun.
	midnight = datetime.combine(today, time.min)
	due = [a for a in assignments if midnight + (a.start_time or timedelta()) <= as_of]
	if not due:
		return []

	employees = [a.employee for a in due]
	punched = set(
		frappe.db.sql_list(
			"""
			SELECT DISTINCT employee FROM `tabEmployee Checkin`
			WHERE employee IN %(employees)s AND DATE(`time`) = %(today)s
			""",
			{"employees": employees, "today": today},
		)
	)
	on_leave = {
		r.employee: r.leave_type
		for r in frappe.db.sql(
			"""
			SELECT employee, leave_type FROM `tabLeave Application`
			WHERE docstatus = 1 AND status = 'Approved'
			  AND employee IN %(employees)s
			  AND from_date <= %(today)s AND to_date >= %(today)s
			""",
			{"employees": employees, "today": today},
			as_dict=True,
		)
	}

	out = []
	for a in due:
		if a.employee in punched or a.employee in on_leave:
			continue
		if _is_holiday(a.holiday_list or company_holiday_list, today):
			continue
		shift_start = midnight + (a.start_time or timedelta())
		out.append(
			frappe._dict(
				employee=a.employee,
				employee_name=a.employee_name,
				department=a.department or "",
				shift=a.shift_type,
				shift_started=shift_start,
				overdue_minutes=int((as_of - shift_start).total_seconds() // 60),
			)
		)
	out.sort(key=lambda r: -r.overdue_minutes)
	return out


def _is_holiday(holiday_list, day):
	if not holiday_list:
		return False
	return bool(
		frappe.db.exists(
			"Holiday", {"parent": holiday_list, "parenttype": "Holiday List", "holiday_date": day}
		)
	)


def get_company_holiday_list(company):
	"""The holiday list that governs this company's working calendar.

	Prefers the list on the company's active Shift Types — that is the one
	HRMS auto-attendance itself honours, and on this site it is the only one
	reliably populated (1 of 299 employees has a personal holiday_list, and
	no company has a default). Falls back to the company default, then to any
	list an employee carries."""
	shift_list = frappe.db.sql_list(
		"""
		SELECT DISTINCT st.holiday_list
		FROM `tabShift Type` st
		WHERE st.holiday_list IS NOT NULL AND st.holiday_list != ''
		  AND st.custom_company = %s
		""",
		company,
	)
	if shift_list:
		return shift_list[0]

	default = frappe.db.get_value("Company", company, "default_holiday_list")
	if default:
		return default

	from_employee = frappe.db.sql_list(
		"""
		SELECT holiday_list FROM `tabEmployee`
		WHERE company = %s AND status = 'Active'
		  AND holiday_list IS NOT NULL AND holiday_list != ''
		LIMIT 1
		""",
		company,
	)
	return from_employee[0] if from_employee else None


def last_working_day(holiday_list, before, max_lookback=14):
	"""The most recent working day strictly before `before`.

	Walks back past holidays so a Sunday digest reports Friday rather than an
	empty Saturday. Gives up after `max_lookback` days and returns the plain
	previous day — a holiday list with a fortnight-long gap is a configuration
	problem, not a reason to send nothing."""
	day = add_days(before, -1)
	for _ in range(max_lookback):
		if not _is_holiday(holiday_list, day):
			return day
		day = add_days(day, -1)
	return add_days(before, -1)


def _device_status(company, as_of):
	devices = frappe.get_all(
		"Biometric Device",
		filters={"company": company, "enabled": 1},
		fields=[
			"name", "device_name", "device_serial", "device_ip",
			"last_contact_time", "last_sync_time", "total_synced",
			"alert_threshold_minutes",
		],
		order_by="device_name asc",
	)
	for d in devices:
		contact = d.last_contact_time or d.last_sync_time
		threshold = cint(d.alert_threshold_minutes) or 120
		if not contact:
			d.silent_minutes = None
			d.verdict = "NEVER REPORTED"
		else:
			d.silent_minutes = int((as_of - get_datetime(contact)).total_seconds() // 60)
			d.verdict = "DOWN" if d.silent_minutes > threshold else "OK"
		d.last_contact = contact
	return devices


# ---------------------------------------------------------------------------
# rendering
# ---------------------------------------------------------------------------

SECTIONS = (
	(
		"missed_punches",
		"Missed Punches",
		["employee", "employee_name", "department", "shift", "punches",
		 "first_punch", "last_punch", "working_hours", "status", "attendance"],
		["Employee", "Name", "Department", "Shift", "Punches",
		 "First Punch", "Last Punch", "Hours", "Attendance Status", "Attendance"],
	),
	(
		"not_seen",
		"Not Seen Today",
		["employee", "employee_name", "department", "shift", "shift_started", "overdue_minutes"],
		["Employee", "Name", "Department", "Shift", "Shift Started", "Minutes Overdue"],
	),
	(
		"devices",
		"Device Status",
		["device_name", "device_serial", "device_ip", "last_contact",
		 "last_sync_time", "silent_minutes", "total_synced", "verdict"],
		["Device", "Serial", "IP", "Last Contact",
		 "Last Sync", "Silent (min)", "Total Synced", "Status"],
	),
)


def _fmt(value):
	if value is None or value == "":
		return "—"
	if isinstance(value, datetime):
		return value.strftime("%Y-%m-%d %H:%M")
	return str(value)


def _html_table(rows, fields, headers, limit=None):
	shown = rows[:limit] if limit else rows
	head = "".join(
		f'<th style="text-align:left;padding:6px 10px;border-bottom:2px solid #d1d5db;'
		f'white-space:nowrap">{h}</th>' for h in headers
	)
	body = []
	for r in shown:
		cells = "".join(
			f'<td style="padding:6px 10px;border-bottom:1px solid #eee;white-space:nowrap">'
			f"{frappe.utils.escape_html(_fmt(r.get(f)))}</td>"
			for f in fields
		)
		body.append(f"<tr>{cells}</tr>")
	more = ""
	if limit and len(rows) > limit:
		more = (
			f'<p style="margin:6px 0 0;color:#6b7280;font-size:12px">'
			f"…and {len(rows) - limit} more — see the attached workbook.</p>"
		)
	return (
		'<table style="border-collapse:collapse;font-size:13px;width:100%">'
		f"<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>{more}"
	)


def render_digest_html(digest):
	missed, unseen, devices = digest["missed_punches"], digest["not_seen"], digest["devices"]
	down = [d for d in devices if d.verdict != "OK"]

	def chip(label, count, bad):
		colour = "#b91c1c" if (bad and count) else "#15803d"
		return (
			f'<span style="display:inline-block;margin:0 14px 0 0">'
			f'<span style="font-size:26px;font-weight:600;color:{colour}">{count}</span>'
			f'<span style="color:#6b7280;font-size:13px"> {label}</span></span>'
		)

	parts = [
		f'<div style="font-family:-apple-system,Segoe UI,Roboto,sans-serif;color:#111">',
		f'<h2 style="margin:0 0 4px">{frappe.utils.escape_html(digest["company"])}</h2>',
		f'<p style="margin:0 0 16px;color:#6b7280;font-size:13px">'
		f'Attendance digest · {digest["as_of"].strftime("%Y-%m-%d %H:%M")}</p>',
		'<p style="margin:0 0 20px">',
		chip(f'missed punches ({digest["for_date"]})', len(missed), True),
		chip("not seen today", len(unseen), True),
		chip("devices down", len(down), True),
		"</p>",
	]

	if not missed and not unseen and not down:
		parts.append(
			'<p style="padding:12px;background:#f0fdf4;border-left:3px solid #16a34a">'
			"Nothing to action. Every punch pairs up, everyone rostered has been seen, "
			"and all devices are reporting.</p>"
		)

	blocks = [
		("Missed Punches", digest["for_date"],
		 "Odd number of punches — a tap is missing, so the hours are wrong. "
		 "Correct these in Attendance Fix before payroll.", missed, 0),
		("Not Seen Today", digest["today"],
		 "Rostered, shift already started, no punch, no approved leave. Chase these now.",
		 unseen, 1),
		("Device Status", None,
		 "A device that stops reporting looks exactly like a workforce that stopped punching.",
		 devices, 2),
	]
	for title, day, blurb, rows, idx in blocks:
		_, _, fields, headers = SECTIONS[idx]
		suffix = f" — {day}" if day else ""
		parts.append(f'<h3 style="margin:24px 0 2px">{title}{suffix}</h3>')
		parts.append(f'<p style="margin:0 0 8px;color:#6b7280;font-size:12px">{blurb}</p>')
		if rows:
			parts.append(_html_table(rows, fields, headers, MAX_ROWS_INLINE))
		else:
			parts.append('<p style="margin:0;color:#15803d;font-size:13px">None.</p>')

	parts.append("</div>")
	return "".join(parts)


def build_workbook(digest):
	"""Three-sheet xlsx. Returns bytes, or None if openpyxl is unavailable."""
	try:
		from openpyxl import Workbook
		from openpyxl.styles import Font
	except ImportError:
		return None

	wb = Workbook()
	wb.remove(wb.active)
	for key, sheet_name, fields, headers in SECTIONS:
		ws = wb.create_sheet(sheet_name)
		ws.append(headers)
		for cell in ws[1]:
			cell.font = Font(bold=True)
		for r in digest[key]:
			ws.append([_fmt(r.get(f)) for f in fields])
		for i, h in enumerate(headers, start=1):
			longest = max([len(h)] + [len(_fmt(r.get(fields[i - 1]))) for r in digest[key]] or [0])
			ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(longest + 2, 40)
		ws.freeze_panes = "A2"

	from io import BytesIO

	buf = BytesIO()
	wb.save(buf)
	return buf.getvalue()


def send_digest_email(company, recipients, digest):
	missed, unseen = len(digest["missed_punches"]), len(digest["not_seen"])
	down = len([d for d in digest["devices"] if d.verdict != "OK"])

	flags = []
	if missed:
		flags.append(f"{missed} missed punch{'es' if missed != 1 else ''}")
	if unseen:
		flags.append(f"{unseen} not seen")
	if down:
		flags.append(f"{down} device{'s' if down != 1 else ''} down")
	summary = ", ".join(flags) if flags else "all clear"

	attachments = []
	workbook = build_workbook(digest)
	if workbook:
		attachments.append({
			"fname": f"attendance-digest-{company[:20].replace(' ', '-')}-{digest['today']}.xlsx",
			"fcontent": workbook,
		})

	frappe.sendmail(
		recipients=recipients,
		subject=_("Attendance digest — {0} — {1}").format(company, summary),
		message=render_digest_html(digest),
		attachments=attachments,
		reference_doctype="Company",
		reference_name=company,
		now=True,
	)


@frappe.whitelist()
def preview_digest(company, for_date=None):
	"""Render the digest for a company without emailing it — lets HR check what
	the 13:00 job will say, and lets us verify a change without waiting a day."""
	frappe.only_for(("System Manager", "HR Manager"))
	return render_digest_html(build_digest(company, for_date=for_date))
