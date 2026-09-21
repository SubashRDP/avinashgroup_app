"""Put a company's payroll on the system from its own monthly salary sheet.

Every company in the group already runs payroll — in Excel. The sheet is the
only record of who is paid what: which allowances each person gets, their grade,
whether they are in the SSF. So onboarding a company means reading its sheet,
not retyping it.

The four sheets use two pay models:

  Initial-basic (NGI, NGN) — one template. Basic Salary, plus HRA as a share of
      a separate Initial Basic, flat gas and education allowances for those
      flagged YES, dearness and other allowances per person, tea by an OLD /
      NEW / NO category, SSF 20% added and 31% deducted.

  Grade-scale (NGG, NGK) — the government scale. Basic plus increments plus
      grade is the salary scale; SSF is 20% / 31% of that scale; then a fixed
      allowance, dearness, other allowance and fuel per person. No HRA, gas or
      education, and no employee codes on the sheet — only names.

A profile per company says which model, where each column is, and the company's
own rates (NGN pays HRA at 27% where NGI pays 25%). `setup_company` builds the
components and the salary structure; `import_sheet` matches the sheet's people to
Employees, writes their pay inputs, and assigns the structure from the start of
the fiscal year. Both are safe to run again.

What varies per person lives on the Employee (the structure reads it in a
formula); what varies per company lives in the structure; what attendance
decides — tea, meals, overtime, late fines — the allowance engine posts each
month. So a rate change is one edit, and a new starter is one Employee.

    bench --site <site> execute avinashgroup_app.payroll.onboarding.onboard \\
        --kwargs "{'key': 'NGN', 'path': '/path/to/NGN.xlsx'}"
"""

import difflib
import re

import frappe
from frappe import _
from frappe.custom.doctype.custom_field.custom_field import create_custom_fields
from frappe.utils import flt

FISCAL_YEAR = "83/84"

INITIAL_BASIC = "initial_basic"
GRADE_SCALE = "grade_scale"

PROFILES = {
	"NGI": {
		"company": "Nepal Gas Udhyog Pvt. Ltd.",
		"model": INITIAL_BASIC,
		"sheet": "Falgun salary",
		"first_row": 5,
		"match": "code",
		"cols": {
			"code": "B", "name": "C", "department": "D", "designation": "F", "attendance": "I",
			"initial_basic": "M", "basic": "N", "hra": "U", "dearness": "V", "other": "W",
			"gas": "Y", "edu": "Z", "tea": "AE", "ssf": "AJ",
		},
		"rates": {"hra": 0.25, "gas": 1690, "edu": 2780, "tea_old": 235, "tea_new": 40, "meal": 75},
	},
	"NGN": {
		"company": "Nepal Gas Udhyog (Narayani) Pvt. Ltd.",
		"model": INITIAL_BASIC,
		"sheet": "Falgun salary",
		"first_row": 5,
		"match": "code",
		"cols": {
			"code": "B", "name": "C", "department": "D", "designation": "E", "attendance": "H",
			"initial_basic": "L", "basic": "M", "hra": "T", "dearness": "U", "other": "V",
			"fuel": "W", "gas": "X", "edu": "Y", "tea": "AD", "ssf": "AH",
		},
		"rates": {"hra": 0.27, "gas": 1690.27, "edu": 1250, "tea_old": 265, "tea_new": 40, "meal": 75},
	},
	"NGG": {
		"company": "Nepal Gas Udhyog (Gandaki) Pvt. Ltd.",
		"model": GRADE_SCALE,
		"sheet": "Falgun salary",
		"first_row": 6,
		"match": "name",
		"cols": {
			"name": "B", "designation": "F", "basic_total": "K", "grade": "O",
			"fixed": "R", "dearness": ("S", "T"), "other": "V", "fuel": "W", "ssf": "AF",
		},
		# 10 holiday OT hours earned 200 of meal on the Falgun sheet: two meals.
		"rates": {"meal": 100},
	},
	"NGK": {
		"company": "Nepal Gas Udhyog (Karnali) Pvt. Ltd.",
		"model": GRADE_SCALE,
		"sheet": "Falgun sal",
		"first_row": 5,
		"match": "name",
		"cols": {
			"name": "B", "designation": "E", "basic_total": "J", "grade": "M",
			"fixed": "P", "dearness": ("Q", "R"), "other": "S", "fuel": "T",
			"maintenance": "U", "ssf": "Y",
		},
		"rates": {},
		# Labour paid by the day: 754 a day for the days worked.
		"wages": {"sheet": "Wages", "first_row": 3, "match": "name", "cols": {"name": "B", "rate": "F"}},
	},
}

EMPLOYEE_FIELDS = [
	{
		"fieldname": "custom_fixed_allowance",
		"label": "Fixed Allowance",
		"fieldtype": "Currency",
		"insert_after": "custom_other_allowance",
		"description": "Grade-scale companies (NGG, NGK): the monthly fixed allowance",
	},
	{
		"fieldname": "custom_maintenance_allowance",
		"label": "Maintenance Allowance",
		"fieldtype": "Currency",
		"insert_after": "custom_fixed_allowance",
	},
]

#: (component, abbr, type) — everything a structure below can hold.
COMPONENTS = (
	("Basic", "B", "Earning"),
	("Dearness Allowance", "DA", "Earning"),
	("Other Allowance", "OA", "Earning"),
	("Fixed Allowance", "FA", "Earning"),
	("Fuel Allowance", "FUEL", "Earning"),
	("Maintenance Allowance", "MNT", "Earning"),
	("House Rent Allowance", "HRA", "Earning"),
	("Gas Allowance", "GAS", "Earning"),
	("Education Allowance", "EDU", "Earning"),
	("SSF Addition", "SSFA", "Earning"),
	# Paid month by month on Payroll Adjustment when there was loading work.
	("Load/Unload Allowance", "LU", "Earning"),
	("SSF", "SSF", "Deduction"),
	("Income Tax", "TAX", "Deduction"),
)


# ───────────────────────────────────────────────────────────── entry point ──


def onboard(key, path):
	"""Set up the company and import its sheet. Returns what happened."""
	setup_company(key)
	result = import_sheet(key, path)
	frappe.db.commit()
	return result


# ────────────────────────────────────────────────────────────────── setup ──


def structure_name(key):
	return f"{key} Staff {FISCAL_YEAR}"


def wage_structure_name(key):
	return f"{key} Daily Wage {FISCAL_YEAR}"


def structure_rows(profile):
	"""(component, formula, depends_on_payment_days) for the company's model.

	Anything a formula reads that is already prorated — Basic — must not be
	prorated again, which is why the two SSF rows are not payment-days rows.
	"""
	r = profile["rates"]
	if profile["model"] == INITIAL_BASIC:
		earnings = [
			("Basic", "base", 1),
			("Dearness Allowance", "custom_dearness_allowance", 1),
			("Other Allowance", "custom_other_allowance", 1),
			("Fuel Allowance", "custom_fuel_allowance", 1),
			("House Rent Allowance", f"custom_initial_basic * {r['hra']} if custom_hra_eligible else 0", 1),
			("Gas Allowance", f"{r['gas']} if custom_gas_allowance else 0", 1),
			("Education Allowance", f"{r['edu']} if custom_education_allowance else 0", 1),
			("SSF Addition", "B * 0.20 if custom_ssf_applicable else 0", 0),
		]
	else:
		earnings = [
			("Basic", "base", 1),
			("Fixed Allowance", "custom_fixed_allowance", 1),
			("Dearness Allowance", "custom_dearness_allowance", 1),
			("Other Allowance", "custom_other_allowance", 1),
			("Fuel Allowance", "custom_fuel_allowance", 1),
		]
		if "maintenance" in profile["cols"]:
			earnings.append(("Maintenance Allowance", "custom_maintenance_allowance", 1))
		earnings.append(("SSF Addition", "B * 0.20 if custom_ssf_applicable else 0", 0))

	deductions = [("SSF", "B * 0.31 if custom_ssf_applicable else 0", 0)]
	return earnings, deductions


def setup_company(key):
	profile = PROFILES[key]
	company = profile["company"]

	create_custom_fields({"Employee": EMPLOYEE_FIELDS}, update=True)
	ensure_payable_account(company)
	ensure_components()
	ensure_allowance_categories(key, profile)

	if profile.get("wages"):
		ensure_wage_structure(key, company)

	name = structure_name(key)
	if frappe.db.exists("Salary Structure", {"name": name, "docstatus": 1}):
		# Assignments hang off a submitted structure, so it is never rebuilt
		# under them; a changed rate is a new structure for the next year.
		return name

	earnings, deductions = structure_rows(profile)
	doc = frappe.get_doc("Salary Structure", name) if frappe.db.exists("Salary Structure", name) else None
	if not doc:
		doc = frappe.new_doc("Salary Structure")
		doc.name = name
	doc.update(
		{
			"company": company,
			"currency": "NPR",
			"payroll_frequency": "Monthly",
			"is_active": "Yes",
			"payment_account": frappe.db.get_value("Company", company, "default_payroll_payable_account"),
			"earnings": [],
			"deductions": [],
		}
	)
	for component, formula, prorate in earnings:
		doc.append("earnings", _row(component, formula, prorate))
	for component, formula, prorate in deductions:
		doc.append("deductions", _row(component, formula, prorate))
	doc.append(
		"deductions",
		{"salary_component": "Income Tax", "variable_based_on_taxable_salary": 1, "depends_on_payment_days": 0},
	)
	doc.flags.ignore_permissions = True
	doc.save()
	doc.submit()
	return doc.name


def ensure_payable_account(company):
	"""Every company's chart has 347301 Salary Payable; Karnali's was never set."""
	if frappe.db.get_value("Company", company, "default_payroll_payable_account"):
		return
	abbr = frappe.db.get_value("Company", company, "abbr")
	account = f"347301 - Salary Payable - {abbr}"
	if frappe.db.exists("Account", account):
		frappe.db.set_value("Company", company, "default_payroll_payable_account", account)


def ensure_wage_structure(key, company):
	"""Daily-wage labour: nothing fixed in the structure. The base is one day's
	pay; the `Daily Wage` component pays it per day present, overtime reads the
	flag, and the tax slab withholds the 1%."""
	name = wage_structure_name(key)
	if frappe.db.exists("Salary Structure", {"name": name, "docstatus": 1}):
		return name
	doc = frappe.new_doc("Salary Structure")
	doc.name = name
	doc.update(
		{
			"company": company,
			"currency": "NPR",
			"payroll_frequency": "Monthly",
			"is_active": "Yes",
			"custom_daily_wage": 1,
			"payment_account": frappe.db.get_value("Company", company, "default_payroll_payable_account"),
		}
	)
	doc.append(
		"deductions",
		{"salary_component": "Income Tax", "variable_based_on_taxable_salary": 1, "depends_on_payment_days": 0},
	)
	doc.flags.ignore_permissions = True
	doc.save()
	doc.submit()
	return doc.name


def _row(component, formula, prorate):
	flags = frappe.db.get_value(
		"Salary Component", component, ["is_tax_applicable", "exempted_from_income_tax"], as_dict=True
	)
	return {
		"salary_component": component,
		"amount_based_on_formula": 1,
		"formula": formula,
		"depends_on_payment_days": prorate,
		# The slip reads these off the row, not the component.
		"is_tax_applicable": flags.is_tax_applicable,
		"exempted_from_income_tax": flags.exempted_from_income_tax,
	}


def ensure_components():
	for name, abbr, kind in COMPONENTS:
		if not frappe.db.exists("Salary Component", name):
			frappe.get_doc(
				{
					"doctype": "Salary Component",
					"salary_component": name,
					"salary_component_abbr": abbr,
					"type": kind,
					"is_tax_applicable": 1 if kind == "Earning" else 0,
					"exempted_from_income_tax": 1 if name == "SSF" else 0,
					"round_to_the_nearest_integer": 0,
				}
			).insert(ignore_permissions=True)


def ensure_allowance_categories(key, profile):
	"""OLD / NEW / NO tea groups and the meal rate, for this company."""
	r = profile["rates"]
	groups = {}
	if "tea_old" in r:
		groups = {"OLD": r["tea_old"], "NEW": r["tea_new"], "NO": 0}
	elif r.get("meal"):
		groups = {"STAFF": None}

	for group, tea in groups.items():
		name = f"{key} {group}"
		if frappe.db.exists("Allowance Category", name):
			continue
		rates = []
		if tea is not None:
			rates.append({"salary_component": "Tea & Conveyance", "rate": tea})
		if r.get("meal"):
			rates.append({"salary_component": "Meal", "rate": r["meal"]})
		frappe.get_doc(
			{
				"doctype": "Allowance Category",
				"category_name": name,
				"company": profile["company"],
				"rates": rates,
			}
		).insert(ignore_permissions=True)


# ───────────────────────────────────────────────────────────────── import ──


def import_sheet(key, path):
	import openpyxl

	profile = PROFILES[key]
	company = profile["company"]
	sheet = openpyxl.load_workbook(path, data_only=True)[profile["sheet"]]
	matcher = EmployeeMatcher(key, company)

	updated, unmatched, assigned, notes = [], [], [], []
	for row in _sheet_rows(sheet, profile):
		employee = matcher.find(row)
		if not employee:
			unmatched.append(row["name"])
			continue

		values, base, note = employee_values(key, profile, row)
		if note:
			notes.append(f"{row['name']}: {note}")
		frappe.db.set_value("Employee", employee, values, update_modified=False)
		updated.append(employee)

		if base and assign(employee, key, company, base):
			assigned.append(employee)

	wages = profile.get("wages")
	wage_assigned, wage_unmatched = [], []
	if wages:
		wage_sheet = openpyxl.load_workbook(path, data_only=True)[wages["sheet"]]
		for row in _sheet_rows(wage_sheet, wages):
			employee = matcher.find(row)
			rate = flt(row.get("rate")) if isinstance(row.get("rate"), (int, float)) else 0
			if not employee or not rate:
				wage_unmatched.append(row["name"])
				continue
			frappe.db.set_value("Employee", employee, "custom_ssf_applicable", 0, update_modified=False)
			if assign(employee, key, company, rate, structure=wage_structure_name(key)):
				wage_assigned.append(employee)

	return {
		"daily_wage_assigned": len(wage_assigned),
		"daily_wage_unmatched": wage_unmatched,
		"company": company,
		"structure": structure_name(key),
		"employees_updated": len(updated),
		"assignments_created": len(assigned),
		"unmatched": unmatched,
		"notes": notes,
	}


def _sheet_rows(sheet, profile):
	from openpyxl.utils import column_index_from_string

	def cell(r, col):
		if isinstance(col, tuple):
			return sum(flt(cell(r, c)) for c in col)
		return sheet.cell(r, column_index_from_string(col)).value

	cols = profile["cols"]
	for r in range(profile["first_row"], sheet.max_row + 1):
		name = cell(r, cols["name"])
		if not isinstance(name, str) or not name.strip():
			continue
		if profile["match"] == "code":
			code = cell(r, cols["code"])
			if not isinstance(code, str) or not re.search(r"\d", code):
				continue
		elif not isinstance(sheet.cell(r, 1).value, (int, float)):
			continue
		yield {field: cell(r, col) for field, col in cols.items()} | {"name": name.strip(), "_row": r}


def employee_values(key, profile, row):
	"""The Employee fields this person's row decides, and their basic."""
	num = lambda v: flt(v) if isinstance(v, (int, float)) else 0.0
	note = None

	if profile["model"] == INITIAL_BASIC:
		base = num(row.get("basic"))
		values = {
			"custom_initial_basic": num(row.get("initial_basic")),
			# The sheet shows the outcome, not the flag: an amount means YES.
			"custom_hra_eligible": 1 if num(row.get("hra")) else 0,
			"custom_gas_allowance": 1 if num(row.get("gas")) else 0,
			"custom_education_allowance": 1 if num(row.get("edu")) else 0,
			"custom_dearness_allowance": num(row.get("dearness")),
			"custom_other_allowance": num(row.get("other")),
			"custom_fuel_allowance": num(row.get("fuel")),
		}
		category = tea_category(key, profile, num(row.get("tea")), num(row.get("attendance")))
		if category:
			values["custom_allowance_category"] = category
	else:
		# Basic + increments + grade, taken from its parts: the sheet's own scale
		# column is prorated for anyone who joined mid-month.
		base = num(row.get("basic_total")) + num(row.get("grade"))
		values = {
			"custom_fixed_allowance": num(row.get("fixed")),
			"custom_dearness_allowance": num(row.get("dearness")),
			"custom_other_allowance": num(row.get("other")),
			"custom_fuel_allowance": num(row.get("fuel")),
			"custom_maintenance_allowance": num(row.get("maintenance")),
		}
		if profile["rates"].get("meal"):
			values["custom_allowance_category"] = f"{key} STAFF"

	values["custom_ssf_applicable"] = 1 if num(row.get("ssf")) else 0

	# Only what this company's sheet actually has: NGI's sheet has no fuel
	# column (its one fuel allowance is paid separately), and writing a 0 for a
	# column that is not there would wipe a figure kept elsewhere.
	source = {
		"custom_initial_basic": "initial_basic", "custom_hra_eligible": "hra",
		"custom_gas_allowance": "gas", "custom_education_allowance": "edu",
		"custom_dearness_allowance": "dearness", "custom_other_allowance": "other",
		"custom_fuel_allowance": "fuel", "custom_fixed_allowance": "fixed",
		"custom_maintenance_allowance": "maintenance", "custom_ssf_applicable": "ssf",
	}
	values = {f: v for f, v in values.items() if f not in source or source[f] in profile["cols"]}

	if not base:
		note = _("no basic on the sheet — not assigned a salary")
	return values, base, note


def tea_category(key, profile, tea, attendance):
	"""OLD / NEW / NO, read back from what the sheet paid for the days worked."""
	r = profile["rates"]
	if not attendance:
		return None
	per_day = tea / attendance
	if not tea:
		return f"{key} NO"
	if abs(per_day - r["tea_old"]) < 1:
		return f"{key} OLD"
	if abs(per_day - r["tea_new"]) < 1:
		return f"{key} NEW"
	return None


def assign(employee, key, company, base, structure=None):
	"""Assign the company structure from the fiscal year's first day, once."""
	fy_start = frappe.db.get_value("Fiscal Year", FISCAL_YEAR, "year_start_date")
	if frappe.db.exists(
		"Salary Structure Assignment",
		{"employee": employee, "docstatus": 1, "from_date": [">=", fy_start]},
	):
		return False

	doc = frappe.new_doc("Salary Structure Assignment")
	doc.update(
		{
			"employee": employee,
			"salary_structure": structure or structure_name(key),
			"from_date": fy_start,
			"company": company,
			"currency": "NPR",
			"base": base,
			"variable": 0,
			# The Payroll Entry finds its employees by this account, so an
			# assignment written without it is silently left out of every run.
			"payroll_payable_account": frappe.db.get_value(
				"Company", company, "default_payroll_payable_account"
			),
			"income_tax_slab": frappe.db.get_value(
				"Income Tax Slab", {"company": company, "docstatus": 1, "disabled": 0}, "name",
				order_by="effective_from desc",
			),
		}
	)
	doc.flags.ignore_permissions = True
	doc.insert()
	doc.submit()
	return True


class EmployeeMatcher:
	"""Sheet row → Employee: by code where the sheet has one, else by name.

	Names on the grade-scale sheets carry titles and stray spacing ("Mr Jiwan
	Chhatkuli", "Rupak Chaudhary "), so names are compared stripped of titles
	and anything but letters, and a near spelling is accepted only when it is
	the single close candidate.
	"""

	TITLES = re.compile(r"^(mr|mrs|ms|miss|dr)\.?\s+", re.I)

	def __init__(self, key, company):
		self.key = key
		self.by_name = {}
		for e in frappe.get_all(
			"Employee", filters={"company": company, "status": "Active"}, fields=["name", "employee_name"]
		):
			self.by_name.setdefault(self.norm(e.employee_name), []).append(e.name)

	@classmethod
	def norm(cls, name):
		return re.sub(r"[^a-z]", "", cls.TITLES.sub("", (name or "").strip()).lower())

	def find(self, row):
		code = row.get("code")
		if code:
			digits = "".join(ch for ch in str(code) if ch.isdigit())
			candidate = f"{self.key}-EMP-{int(digits):05d}"
			if frappe.db.exists("Employee", candidate):
				return candidate

		key = self.norm(row["name"])
		exact = self.by_name.get(key)
		if exact and len(exact) == 1:
			return exact[0]

		close = difflib.get_close_matches(key, list(self.by_name), n=2, cutoff=0.88)
		if len(close) == 1 and len(self.by_name[close[0]]) == 1:
			return self.by_name[close[0]][0]
		return None


# ─────────────────────────────────────────────────────────────── accounts ──

#: A PROPOSAL for the accountant, not applied by `onboard`. Every company's chart
#: carries the same codes, so one mapping serves all seven. Proven end to end on
#: nepalgas (inside a rolled-back transaction): 87 NGI slips submitted and the
#: accrual journal balanced, crediting Salary Payable with exactly the net pay.
#:
#: Open question for the accountant: the chart has three salary expense accounts
#: (547101 O/O office, 547102 S/D sales & distribution, 547103 F/P filling
#: plant) but no matching cost centres, and a component maps to one account. So
#: this posts every earning to 547101 until cost centres exist to split it.
PROPOSED_ACCOUNTS = {
	"__earnings__": "547101 - Salary Expenses - O/O",
	"SSF": "347302 - SSF Payable",
	"Additional SSF": "347302 - SSF Payable",
	"Income Tax": "348102 - 11112 TDS-Remuneration Income Tax",
	"Salary Advance": "148003 - Staff Advance",
	# No fines-income account exists; a fine reduces what the salary cost.
	"Late Fine": "547101 - Salary Expenses - O/O",
}


def map_component_accounts(company):
	"""Give every salary component its GL account for one company.

	Without this the payroll journal cannot be written and submitting the
	month's salary slips fails. Run only once the accountant has agreed the
	mapping above. Existing rows are left as they are.
	"""
	abbr = frappe.db.get_value("Company", company, "abbr")
	done = []
	for component in frappe.get_all("Salary Component", fields=["name", "type"]):
		base = PROPOSED_ACCOUNTS.get(component.name) or (
			PROPOSED_ACCOUNTS["__earnings__"] if component.type == "Earning" else None
		)
		if not base:
			continue
		account = f"{base} - {abbr}"
		if not frappe.db.exists("Account", account):
			continue
		if frappe.db.exists(
			"Salary Component Account", {"parent": component.name, "company": company}
		):
			continue
		doc = frappe.get_doc("Salary Component", component.name)
		doc.append("accounts", {"company": company, "account": account})
		doc.flags.ignore_permissions = True
		doc.save()
		done.append((component.name, account))
	return done
