# Copyright (c) 2013, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt

import json
import os
from collections import defaultdict

import frappe
from frappe import _
from frappe.utils import cint, flt, fmt_money


def _as_list(value):
	"""Normalize a MultiSelectList/Link filter value (list, JSON string, or single) to a list."""
	if not value:
		return []
	if isinstance(value, str):
		value = value.strip()
		if value.startswith("["):
			try:
				value = json.loads(value)
			except Exception:
				return [value]
		else:
			return [value]
	if isinstance(value, (list, tuple, set)):
		return [v for v in value if v]
	return [value]


def fmt_qty(value):
	"""A quantity for print and email: grouped, with only the decimals it needs
	(1, 1.5, 1,250). The on-screen report formats its own."""
	if value is None or value == "":
		return ""
	return "{:,.3f}".format(flt(value)).rstrip("0").rstrip(".")


def _sq_filter_scope(company=None, purchase_order=None, material_request=None,
					 supplier_quotation=None, supplier=None, item_code=None):
	"""Shared WHERE fragments (aliases sq / sqi) limiting quotation rows to the
	report's current filter context. The filter-option endpoints below all build
	on this so every dropdown only offers values that can actually appear in the
	report for the chosen Purchase Order / Material Request / company."""
	conditions = ["sqi.parent = sq.name", "sqi.docstatus < 2"]
	values = {}

	if company:
		conditions.append("sq.company = %(company)s")
		values["company"] = company

	# Arrives JSON-encoded from the dropdowns ('["PO-1"]', '[]'), so normalise
	# before testing it - a bare '[]' string would otherwise read as "PO chosen".
	purchase_orders = _as_list(purchase_order)
	material_requests = _as_list(material_request)
	if purchase_orders:
		material_requests += get_material_requests_from_purchase_order(purchase_orders)
	material_requests = [mr for mr in dict.fromkeys(material_requests) if mr]
	if material_requests:
		conditions.append("sqi.material_request IN %(material_requests)s")
		values["material_requests"] = tuple(material_requests)
	elif purchase_orders:
		# PO chosen but untraceable to any MR -> no quotation is aligned to it;
		# offer nothing rather than everything (mirrors get_data).
		conditions.append("1 = 0")

	quotations = _as_list(supplier_quotation)
	if quotations:
		conditions.append("sq.name IN %(quotations)s")
		values["quotations"] = tuple(quotations)

	suppliers = _as_list(supplier)
	if suppliers:
		conditions.append("sq.supplier IN %(suppliers)s")
		values["suppliers"] = tuple(suppliers)

	if item_code:
		conditions.append("sqi.item_code = %(item_code)s")
		values["item_code"] = item_code

	return conditions, values


@frappe.whitelist()
def get_filter_suppliers(company=None, purchase_order=None, material_request=None,
						 supplier_quotation=None, item_code=None, txt=None):
	"""Supplier options: only suppliers with a quotation in the current scope."""
	conditions, values = _sq_filter_scope(
		company=company, purchase_order=purchase_order, material_request=material_request,
		supplier_quotation=supplier_quotation, item_code=item_code,
	)
	conditions.append("(sq.supplier LIKE %(txt)s OR sq.supplier_name LIKE %(txt)s)")
	values["txt"] = f"%{(txt or '').strip()}%"

	return frappe.db.sql(
		f"""
		SELECT DISTINCT sq.supplier AS value, sq.supplier_name AS description
		FROM `tabSupplier Quotation` sq, `tabSupplier Quotation Item` sqi
		WHERE {" AND ".join(conditions)}
		ORDER BY sq.supplier_name
		LIMIT 50
		""",
		values,
		as_dict=True,
	)


@frappe.whitelist()
def get_filter_purchase_orders(company=None, material_request=None, purchase_order=None, txt=None):
	"""Purchase Order options: only orders raised from a Material Request that is in
	the comparison. That set is the Material Request filter plus the MRs of the
	orders already picked - so opened from one PO, the list is that PO and its
	siblings (the same MR ordered from other suppliers). With nothing picked yet it
	is any PO whose MR has a quotation to compare; a PO with no MR can never
	produce a comparison, so it is never offered."""
	material_requests = _as_list(material_request)
	purchase_orders = _as_list(purchase_order)
	if purchase_orders:
		material_requests += get_material_requests_from_purchase_order(purchase_orders)
	material_requests = [mr for mr in dict.fromkeys(material_requests) if mr]

	conditions = ["poi.parent = po.name", "po.docstatus < 2", "poi.material_request IS NOT NULL"]
	values = {"txt": f"%{(txt or '').strip()}%"}
	if company:
		conditions.append("po.company = %(company)s")
		values["company"] = company
	if material_requests:
		conditions.append("poi.material_request IN %(material_requests)s")
		values["material_requests"] = tuple(material_requests)
	else:
		conditions.append(
			"""EXISTS (
				SELECT 1 FROM `tabSupplier Quotation Item` sqi
				WHERE sqi.material_request = poi.material_request AND sqi.docstatus < 2
			)"""
		)
	conditions.append("(po.name LIKE %(txt)s OR po.supplier_name LIKE %(txt)s)")

	return frappe.db.sql(
		f"""
		SELECT po.name AS value, po.supplier_name AS description
		FROM `tabPurchase Order` po, `tabPurchase Order Item` poi
		WHERE {" AND ".join(conditions)}
		GROUP BY po.name
		ORDER BY po.transaction_date DESC, po.name DESC
		LIMIT 50
		""",
		values,
		as_dict=True,
	)


@frappe.whitelist()
@frappe.validate_and_sanitize_search_inputs
def get_filter_items(doctype, txt, searchfield, start, page_len, filters):
	"""Item link-query: only items quoted in the current scope."""
	filters = filters or {}
	conditions, values = _sq_filter_scope(
		company=filters.get("company"),
		purchase_order=filters.get("purchase_order"),
		material_request=filters.get("material_request"),
		supplier_quotation=filters.get("supplier_quotation"),
		supplier=filters.get("supplier"),
	)
	conditions.append("(sqi.item_code LIKE %(txt)s OR sqi.item_name LIKE %(txt)s)")
	values.update({"txt": f"%{(txt or '').strip()}%", "start": cint(start), "page_len": cint(page_len) or 20})

	return frappe.db.sql(
		f"""
		SELECT DISTINCT sqi.item_code, sqi.item_name
		FROM `tabSupplier Quotation` sq, `tabSupplier Quotation Item` sqi
		WHERE {" AND ".join(conditions)}
		ORDER BY sqi.item_name
		LIMIT %(start)s, %(page_len)s
		""",
		values,
	)


@frappe.whitelist()
def get_supplier_quotations(company=None, purchase_order=None, material_request=None,
							supplier=None, item_code=None, txt=None):
	"""Supplier Quotation options: only quotations in the current scope."""
	conditions, values = _sq_filter_scope(
		company=company, purchase_order=purchase_order, material_request=material_request,
		supplier=supplier, item_code=item_code,
	)
	conditions.append("(sq.name LIKE %(txt)s OR sq.supplier_name LIKE %(txt)s)")
	values["txt"] = f"%{(txt or '').strip()}%"

	return frappe.db.sql(
		f"""
		SELECT DISTINCT sq.name AS value, sq.supplier_name AS description
		FROM `tabSupplier Quotation` sq, `tabSupplier Quotation Item` sqi
		WHERE {" AND ".join(conditions)}
		ORDER BY sq.transaction_date DESC
		LIMIT 50
		""",
		values,
		as_dict=True,
	)


@frappe.whitelist()
@frappe.validate_and_sanitize_search_inputs
def get_filter_material_requests(doctype, txt, searchfield, start, page_len, filters):
	"""Material Request link-query: only MRs with a quotation in the current scope
	(company / Purchase Order(s) / Supplier / Supplier Quotation / Item). With a PO
	picked that is the PO's own MR(s)."""
	filters = filters or {}
	conditions, values = _sq_filter_scope(
		company=filters.get("company"),
		purchase_order=filters.get("purchase_order"),
		supplier_quotation=filters.get("supplier_quotation"),
		supplier=filters.get("supplier"),
		item_code=filters.get("item_code"),
	)
	conditions += ["mr.name = sqi.material_request", "mr.docstatus < 2", "mr.name LIKE %(txt)s"]
	values.update({"txt": f"%{(txt or '').strip()}%", "start": cint(start), "page_len": cint(page_len) or 20})

	return frappe.db.sql(
		f"""
		SELECT mr.name, mr.transaction_date
		FROM `tabSupplier Quotation` sq, `tabSupplier Quotation Item` sqi, `tabMaterial Request` mr
		WHERE {" AND ".join(conditions)}
		GROUP BY mr.name
		ORDER BY mr.transaction_date DESC, mr.name DESC
		LIMIT %(start)s, %(page_len)s
		""",
		values,
	)


def execute(filters=None):
	if not filters:
		return [], []

	# Get raw data
	supplier_quotation_data = get_data(filters)

	# Prepare pivoted data and get list of suppliers
	(
		data,
		quotations,
		quotation_display_name,
		column_key,
		quotation_pos,
		quotations_with_narration,
		po_key,
	) = prepare_pivoted_data(supplier_quotation_data, filters)

	# Generate columns dynamically: a group per quotation, and inside it the Quoted
	# sub-group plus one sub-group per Purchase Order made from that quotation
	columns = get_columns(
		filters,
		quotations,
		quotation_display_name,
		column_key,
		quotation_pos,
		quotations_with_narration,
		po_key,
	)

	message = get_message()

	return columns, data, message, None


def get_data(filters):
	"""Fetch supplier quotation data from database"""
	sq = frappe.qb.DocType("Supplier Quotation")
	sq_item = frappe.qb.DocType("Supplier Quotation Item")

	query = (
		frappe.qb.from_(sq_item)
		.from_(sq)
		.select(
			sq_item.parent,
			sq_item.item_code,
			sq_item.item_name,
			sq_item.custom_narration,
			sq_item.material_request_item,
			sq_item.qty,
			sq.currency,
			sq_item.stock_qty,
			sq_item.amount,
			sq_item.rate,
			sq_item.base_rate,
			sq_item.base_amount,
			sq.price_list_currency,
			sq_item.uom,
			sq_item.stock_uom,
			sq_item.request_for_quotation,
			sq_item.lead_time_days,
			sq.supplier.as_("supplier_id"),
			sq.supplier_name,
			sq.valid_till,
			sq.custom_specification,
			sq.custom_warrenty,
			sq.custom_payment_terms,
			sq.custom_delivery_period,
			sq.transaction_date,
			sq.taxes_and_charges,
			sq.discount_amount,
			sq.apply_discount_on,
			sq.additional_discount_percentage,
			sq.grand_total,
			sq.base_grand_total,
			sq.total,
			sq.base_total,
			sq.total_taxes_and_charges,
			sq.base_total_taxes_and_charges,
		)
		.orderby(sq_item.item_code, sq.supplier)
	)

	query = query.where(
		(sq_item.parent == sq.name)
		& (sq_item.docstatus < 2)
		& (sq.company == filters.get("company"))
	)

	# The date window is a convenience for browsing the report on its own, so it is
	# applied only when given. Opened from a Purchase Order it is deliberately left
	# empty: that PO's Material Request(s) are already an exact scope, and a window
	# could only ever hide quotations from it - a quotation is always raised before
	# the order it leads to, so a range starting at the PO's own date excludes every
	# one of them.
	if filters.get("from_date"):
		query = query.where(sq.transaction_date >= filters.get("from_date"))
	if filters.get("to_date"):
		query = query.where(sq.transaction_date <= filters.get("to_date"))

	# Source-document filter. A Supplier Quotation Item links directly to the
	# Material Request it was raised from. A Purchase Order carries the Material
	# Request on its item lines rather than on the quotation, so we resolve the PO
	# to its Material Request(s) and match on the same column. When both filters are
	# set we take the union — "show quotations for any of these source documents".
	# Several Purchase Orders may be chosen; each contributes its own MR(s).
	purchase_orders = _as_list(filters.get("purchase_order"))
	material_requests = _as_list(filters.get("material_request"))
	if purchase_orders:
		material_requests += get_material_requests_from_purchase_order(purchase_orders)
	material_requests = [mr for mr in dict.fromkeys(material_requests) if mr]
	if material_requests:
		query = query.where(sq_item.material_request.isin(material_requests))
	elif purchase_orders:
		# Purchase Orders were chosen but none of their item lines carries a
		# material_request link (created directly, not from an MR). Nothing is
		# "aligned to these POs", so show nothing - falling through would list
		# every quotation in the window.
		return []

	if filters.get("item_code"):
		query = query.where(sq_item.item_code == filters.get("item_code"))

	if filters.get("supplier_quotation"):
		query = query.where(sq_item.parent.isin(filters.get("supplier_quotation")))

	if filters.get("supplier"):
		query = query.where(sq.supplier.isin(filters.get("supplier")))

	if filters.get("preferred_quotation"):
		query = query.where(sq.custom_preferred_quotation == 1)

	supplier_quotation_data = query.run(as_dict=True)

	return supplier_quotation_data


def prepare_pivoted_data(supplier_quotation_data, filters):
	"""
	Transform row-based data into pivoted format with totals:
	- Items as rows
	- One column block per Supplier Quotation. A supplier that quoted more than
	  once for the same request gets one block per quotation, numbered
	  "Supplier 1", "Supplier 2", ... so each quotation is compared on its own
	  terms instead of being collapsed into a single column
	- Subtotal, discount, tax, and invoice total rows
	"""

	# ============================================
	# CONFIGURABLE: Change this to use different price field
	# Options: 'base_rate', 'base_amount', 'rate', 'amount'
	# ============================================
	if not supplier_quotation_data:
		return [], [], {}, {}, {}, set(), {}

	price_field = filters.get("price_field", "base_amount")
	rate_field = "base_rate" if price_field.startswith("base_") else "rate"

	float_precision = cint(frappe.db.get_default("float_precision")) or 2

	# Data structures for pivot. Each is keyed by Supplier Quotation name - one
	# quotation is one column block in the report.
	item_quotation_map = defaultdict(lambda: defaultdict(dict))
	quotation_map = {}  # Store quotation-level data per quotation
	quotation_supplier = {}  # quotation -> (supplier id, supplier display name)
	all_items = []  # Maintain order
	item_meta = {}  # Store item metadata
	seen_items = set()
	item_mr_lines = defaultdict(set)  # item -> Material Request Item rows its quotation lines answer

	# Process each quotation line
	for row in supplier_quotation_data:
		item_code = row.get("item_code")
		quotation_name = row.get("parent")
		if not quotation_name:
			continue
		quotation_supplier[quotation_name] = (
			row.get("supplier_id"),
			row.get("supplier_name") or row.get("supplier_id"),
		)

		# Get price value based on configured field
		price_value = flt(row.get(price_field), float_precision)
		rate_value = flt(row.get(rate_field), float_precision)

		if item_code not in seen_items:
			all_items.append(item_code)
			seen_items.add(item_code)

		# Store item metadata (use first occurrence)
		if item_code not in item_meta:
			item_meta[item_code] = {
				"item_name": row.get("item_name"),
				"uom": row.get("uom"),
			}

		# The Material Request line this quotation line answers - its qty is what
		# was asked for, which the MR Qty column shows (see requested_qty below).
		if row.get("material_request_item"):
			item_mr_lines[item_code].add(row.get("material_request_item"))

		# Narration (custom_narration, Small Text) is the quotation's own
		# free-text note on its line for this item, so it is stored per
		# (item, quotation) alongside the price.
		narration_value = (row.get("custom_narration") or "").strip()

		# Store this quotation's price for the item. A quotation can list the
		# same item on more than one line - keep the cheapest of those lines.
		existing = item_quotation_map[item_code].get(quotation_name)
		if not existing or price_value < existing["price"]:
			item_quotation_map[item_code][quotation_name] = {
				"price": price_value,
				"rate": rate_value,
				"qty": row.get("qty"),
				"narration": narration_value,
			}

		# Store quotation-level data (discount, taxes) per quotation
		if quotation_name not in quotation_map:
			quotation_map[quotation_name] = {
				"discount_amount": flt(row.get("discount_amount"), float_precision),
				"additional_discount_percentage": flt(row.get("additional_discount_percentage"), float_precision),
				"total": flt(row.get("base_total"), float_precision),
				"total_taxes": flt(row.get("base_total_taxes_and_charges"), float_precision),
				"grand_total": flt(row.get("base_grand_total"), float_precision),
				"apply_discount_on": row.get("apply_discount_on"),
				# Quotation-level commercial terms, rendered as free-text rows below
				# the invoice total (one row per attribute, value per quotation).
				"specification": row.get("custom_specification"),
				"warranty": row.get("custom_warrenty"),
				"payment_terms": row.get("custom_payment_terms"),
				"delivery_period": row.get("custom_delivery_period"),
			}

	# MR Qty = what the Material Request asked for: the qty of every MR line the
	# quotations answer, per item (an item can sit on more than one MR line or MR).
	# Each quotation's own qty is shown in its block as Quoted. Blank for an item
	# whose quotation lines carry no MR link.
	mr_lines = {name for names in item_mr_lines.values() for name in names}
	mr_line_qty = dict(frappe.db.sql(
		"SELECT name, qty FROM `tabMaterial Request Item` WHERE name IN %(lines)s",
		{"lines": tuple(mr_lines)},
	)) if mr_lines else {}
	requested_qty = {
		item_code: sum(flt(mr_line_qty.get(name)) for name in names)
		for item_code, names in item_mr_lines.items()
	}

	# Column order: suppliers alphabetically, with a supplier's own quotations
	# adjacent and in name order - which is the order they are numbered in below.
	quotations = sorted(quotation_supplier, key=lambda q: (quotation_supplier[q][1] or "", q))

	# A supplier holding more than one quotation here gets its blocks numbered
	# ("Acme 1", "Acme 2"). The display name is also what the merged supplier
	# header groups consecutive columns on, so it has to differ between the
	# blocks - otherwise two quotations would be drawn as one supplier.
	quotations_per_supplier = defaultdict(list)
	for quotation in quotations:
		quotations_per_supplier[quotation_supplier[quotation][0]].append(quotation)

	quotation_display_name = {}
	for supplier_quotations in quotations_per_supplier.values():
		numbered = len(supplier_quotations) > 1
		for position, quotation in enumerate(supplier_quotations, start=1):
			display = quotation_supplier[quotation][1]
			quotation_display_name[quotation] = f"{display} {position}" if numbered else display

	# Fieldname stem per block. frappe.scrub leaves "/" in place and the naming
	# series carries the fiscal year (NGI-SQ-83/84-00010), which has no business
	# in a fieldname; the guard keeps two names from collapsing onto one stem.
	column_key = {}
	used_keys = set()
	for position, quotation in enumerate(quotations, start=1):
		key = frappe.scrub(quotation).replace("/", "_")
		if key in used_keys:
			key = f"{key}_{position}"
		used_keys.add(key)
		column_key[quotation] = key

	# Quotations with at least one non-blank narration anywhere in the comparison -
	# like Ordered, the Narration column is dropped entirely for a quotation with
	# none, rather than showing an always-blank column.
	quotations_with_narration = {
		quotation
		for item_quotations in item_quotation_map.values()
		for quotation, line in item_quotations.items()
		if (line.get("narration") or "").strip()
	}

	# ------------------------------------------------------------------
	# The Purchase Orders made from each quotation. Purchase Order Item.
	# supplier_quotation records the quotation each PO line came from, so each
	# quotation gets one sub-group per PO raised from it, showing that PO's own
	# qty, rate and amount per item beside what was quoted. Every such PO is
	# shown whatever the Purchase Order filter says: the filter picks which
	# Material Requests are compared and stars its own POs, it does not hide
	# the sibling orders.
	# ------------------------------------------------------------------
	po_lines = {}                      # (item_code, quotation, po) -> {"qty", "amount"}
	quotation_pos = defaultdict(list)  # quotation -> its POs, in name order
	if quotations:
		for pr in frappe.db.sql(
			"""
			SELECT poi.item_code AS item_code, poi.supplier_quotation AS sq, po.name AS po,
			       SUM(poi.qty) AS qty, SUM(poi.base_amount) AS amount
			FROM `tabPurchase Order Item` poi, `tabPurchase Order` po
			WHERE poi.parent = po.name AND po.docstatus < 2 AND poi.supplier_quotation IN %(sqs)s
			GROUP BY poi.item_code, poi.supplier_quotation, po.name
			ORDER BY po.name
			""",
			{"sqs": tuple(quotations)},
			as_dict=True,
		):
			po_lines[(pr.item_code, pr.sq, pr.po)] = {
				"qty": flt(pr.qty),
				"amount": flt(pr.amount, float_precision),
			}
			if pr.po not in quotation_pos[pr.sq]:
				quotation_pos[pr.sq].append(pr.po)

	# Fieldname stem per PO, cleaned like the quotation stems (names carry "/").
	po_key = {po: frappe.scrub(po).replace("/", "_") for pos in quotation_pos.values() for po in pos}
	po_totals = defaultdict(float)  # (quotation, po) -> amount of this quotation's items on that PO

	# Build output rows
	data = []
	sn = 1
	quotation_totals = {quotation: 0 for quotation in quotations}

	# Item rows
	for item_code in all_items:
		row = {
			"sn": sn,
			"item_code": item_code,
			"item_name": item_meta[item_code]["item_name"],
			"qty": requested_qty.get(item_code),
			"is_data_row": 1,  # Mark as data row for styling
		}

		sn += 1

		# Add each quotation's rate + amount + narration as columns
		for quotation in quotations:
			col_fieldname = column_key[quotation]

			if quotation in item_quotation_map[item_code]:
				line = item_quotation_map[item_code][quotation]
				price = line["price"]
				row[col_fieldname + "_qty"] = flt(line["qty"])
				row[col_fieldname + "_rate"] = line["rate"]
				row[col_fieldname] = price
				row[col_fieldname + "_narration"] = line.get("narration") or ""
				quotation_totals[quotation] += price
			else:
				row[col_fieldname + "_qty"] = None
				row[col_fieldname + "_rate"] = None
				row[col_fieldname] = None
				row[col_fieldname + "_narration"] = None

			# Each Purchase Order made from this quotation: that PO's own qty, rate
			# and amount for the item (blank when the PO did not take it).
			for po in quotation_pos.get(quotation, []):
				stem = f"{col_fieldname}__{po_key[po]}"
				line = po_lines.get((item_code, quotation, po))
				if line and line["qty"]:
					row[stem + "_poqty"] = line["qty"]
					row[stem + "_porate"] = flt(line["amount"] / line["qty"], float_precision)
					row[stem + "_poamt"] = line["amount"]
					po_totals[(quotation, po)] += line["amount"]
				else:
					row[stem + "_poqty"] = row[stem + "_porate"] = row[stem + "_poamt"] = None

		data.append(row)
	
	# ============================================
	# SUMMARY ROWS
	# ============================================
	
	def summary_row(label, **flags):
		return {"sn": None, "item_code": None, "item_name": None, "qty": label, **flags}

	def has_value(row):
		"""True if any quotation column in the row is non-zero."""
		return any(flt(row[column_key[q]]) for q in quotations)

	# Total row (Net Total)
	total_row = summary_row("Total", is_total_row=1)
	for quotation in quotations:
		total_row[column_key[quotation]] = quotation_totals[quotation]
		# and each PO's own total for this quotation's items
		for po in quotation_pos.get(quotation, []):
			total_row[f"{column_key[quotation]}__{po_key[po]}_poamt"] = po_totals[(quotation, po)]
	data.append(total_row)

	# Discount on Net Total row
	discount_net_row = summary_row("Less: Discount (on Net Total)", is_summary_row=1)
	for quotation in quotations:
		col_fieldname = column_key[quotation]
		discount_amount = 0

		if quotation in quotation_map:
			sq_data = quotation_map[quotation]
			apply_on = sq_data.get("apply_discount_on")

			# Only show discount here if it's applied on Net Total
			if apply_on == "Net Total":
				discount_amount = flt(sq_data.get("discount_amount", 0), float_precision)

		discount_net_row[col_fieldname] = discount_amount
	if has_value(discount_net_row):
		data.append(discount_net_row)

	# Taxable Amount row (after net discount) - always shown
	taxable_row = summary_row("Taxable Amount", is_summary_row=1)
	for quotation in quotations:
		col_fieldname = column_key[quotation]
		total = total_row[col_fieldname]
		discount = discount_net_row[col_fieldname]
		taxable_row[col_fieldname] = total - discount
	data.append(taxable_row)

	# VAT/Tax row
	vat_row = summary_row("Add: VAT", is_summary_row=1)
	for quotation in quotations:
		col_fieldname = column_key[quotation]
		if quotation in quotation_map:
			vat_row[col_fieldname] = flt(quotation_map[quotation].get("total_taxes", 0), float_precision)
		else:
			vat_row[col_fieldname] = 0
	if has_value(vat_row):
		data.append(vat_row)

	# Discount on Grand Total row
	discount_grand_row = summary_row("Less: Discount (on Grand Total)", is_summary_row=1)
	for quotation in quotations:
		col_fieldname = column_key[quotation]
		discount_amount = 0

		if quotation in quotation_map:
			sq_data = quotation_map[quotation]
			apply_on = sq_data.get("apply_discount_on")

			# Only show discount here if it's applied on Grand Total
			if apply_on == "Grand Total":
				discount_amount = flt(sq_data.get("discount_amount", 0), float_precision)

		discount_grand_row[col_fieldname] = discount_amount
	if has_value(discount_grand_row):
		data.append(discount_grand_row)

	# Invoice Amount row (Grand Total) - use DB value directly, always shown
	invoice_row = summary_row("Invoice Amount", is_invoice_row=1)
	for quotation in quotations:
		col_fieldname = column_key[quotation]
		if quotation in quotation_map:
			# Use grand_total directly from database
			invoice_row[col_fieldname] = flt(quotation_map[quotation].get("grand_total", 0), float_precision)
		else:
			# Fallback calculation if no quotation data
			taxable = taxable_row[col_fieldname]
			vat = vat_row[col_fieldname]
			discount_grand = discount_grand_row[col_fieldname]
			invoice_row[col_fieldname] = taxable + vat - discount_grand
	data.append(invoice_row)

	# ============================================
	# COMMERCIAL-TERMS ROWS
	# One free-text row per quotation-level attribute (Specification, Warranty,
	# Payment Terms, Delivery Period), the value shown under each quotation. A row
	# is dropped entirely when no quotation in the comparison carries that field.
	# Marked is_term_row so the client formatter / PDF / email skip currency
	# formatting and the per-unit Rate column for these lines.
	# ============================================
	term_rows = [
		(_("Specification"), "specification"),
		(_("Warranty"), "warranty"),
		(_("Payment Terms"), "payment_terms"),
		(_("Delivery Period"), "delivery_period"),
	]
	for label, key in term_rows:
		term_row = summary_row(label, is_term_row=1)
		for quotation in quotations:
			value = ""
			if quotation in quotation_map:
				value = (quotation_map[quotation].get(key) or "").strip()
			term_row[column_key[quotation]] = value
		if any(term_row[column_key[q]] for q in quotations):
			data.append(term_row)

	return (
		data,
		quotations,
		quotation_display_name,
		column_key,
		dict(quotation_pos),
		quotations_with_narration,
		po_key,
	)


def get_columns(
	filters,
	quotations,
	quotation_display_name=None,
	column_key=None,
	quotation_pos=None,
	quotations_with_narration=None,
	po_key=None,
):
	"""
	Generate columns dynamically, three header levels deep like a hand-made sheet:
	- Fixed columns for SN, item name and MR Qty
	- A group per Supplier Quotation (`supplier_group`, `sq_link`)
	- Inside it the "Quoted" sub-group (the quotation's Qty / Rate / Amount /
	  Narration), then one sub-group per Purchase Order made from that quotation
	  (the PO's own Qty / Rate / Amount), labelled with the PO number (`sub_group`,
	  `po_link`). A PO picked in the Purchase Order filter is starred.
	"""
	# Fixed columns
	columns = [
		{
			"fieldname": "sn",
			"label": _("SN"),
			"fieldtype": "Int",
			"width": 40,
		},
		{
			"fieldname": "item_name",
			"label": _("Item Name"),
			"fieldtype": "Data",
			"width": 140,
		},
		{
			"fieldname": "qty",
			# What the Material Request asked for; each quotation's own qty is in its block
			"label": _("MR Qty"),
			"fieldtype": "Data",  # Data type to allow "Total" text
			"width": 60,
		},
	]

	column_key = column_key or {}
	quotation_pos = quotation_pos or {}
	quotations_with_narration = quotations_with_narration or set()
	po_key = po_key or {}
	starred = set(_as_list((filters or {}).get("purchase_order")))

	for quotation in quotations:
		col_fieldname = column_key.get(quotation) or frappe.scrub(quotation).replace("/", "_")
		block = {
			"sq_link": quotation,
			"supplier_group": (quotation_display_name or {}).get(quotation, quotation),
		}

		# What the quotation offered. Narration is dropped for a quotation that has
		# none anywhere in the comparison, rather than showing an always-blank column.
		quoted = {**block, "sub_group": _("Quoted")}
		columns += [
			{**quoted, "fieldname": col_fieldname + "_qty", "label": _("Qty"),
			 "fieldtype": "Float", "width": 70},
			{**quoted, "fieldname": col_fieldname + "_rate", "label": _("Rate"),
			 "fieldtype": "Currency", "options": "Company:company:default_currency", "width": 85},
			{**quoted, "fieldname": col_fieldname, "label": _("Amount"),
			 "fieldtype": "Currency", "options": "Company:company:default_currency", "width": 100},
		]
		if quotation in quotations_with_narration:
			columns.append({**quoted, "fieldname": col_fieldname + "_narration",
							"label": _("Narration"), "fieldtype": "Data", "width": 90})

		# What each Purchase Order made from it actually ordered.
		for po in quotation_pos.get(quotation, []):
			stem = f"{col_fieldname}__{po_key[po]}"
			ordered = {**block, "sub_group": ("★ " if po in starred else "") + po, "po_link": po}
			columns += [
				{**ordered, "fieldname": stem + "_poqty", "label": _("Qty"),
				 "fieldtype": "Float", "width": 70},
				{**ordered, "fieldname": stem + "_porate", "label": _("Rate"),
				 "fieldtype": "Currency", "options": "Company:company:default_currency", "width": 85},
				{**ordered, "fieldname": stem + "_poamt", "label": _("Amount"),
				 "fieldtype": "Currency", "options": "Company:company:default_currency", "width": 100},
			]

	return columns


def get_message():
	"""Return report message/legend"""
	return f"""<span class="indicator blue">
		{_("Comparison sheet showing item-wise quotations from all suppliers")}
		</span>
		<br>
		<span class="indicator">
		{_("Total includes all item amounts before discount")}
		</span>
		<br>
		<span class="indicator green">
		{_("Each quotation shows what was Quoted, then every Purchase Order made from it (★ = the order you opened). Click a heading to open it")}
		</span>
		<br>
		<span class="indicator orange">
		{_("A supplier that quoted more than once gets one numbered column set per quotation")}
		</span>"""


# The alternating supplier tints the on-screen report paints (kept in step with
# BAND_COLORS in the report's client script) so print and Excel read the same way.
SUPPLIER_BAND_COLORS = ["#eef3ff", "#fff8ec", "#eefaf1", "#fdeef4"]

# Suffix -> kind, for classifying a quotation block's columns. The PO suffixes are
# checked first; anything unmatched is the quotation's Amount column, which carries
# the bare scrubbed quotation name.
_SUPPLIER_FIELD_KINDS = (
	("_poqty", "poqty"),
	("_porate", "porate"),
	("_poamt", "poamt"),
	("_narration", "narration"),
	("_rate", "rate"),
	("_qty", "qty"),
)


def _supplier_groups(columns):
	"""The column blocks exactly as the on-screen report draws them, three header
	levels deep: a group per Supplier Quotation (consecutive columns sharing an
	`sq_link`), inside it a sub-group per `sub_group` - "Quoted", then one per
	Purchase Order made from that quotation - and inside those the columns. Print,
	Excel and the approval email all walk this instead of restating the report's
	column rules. Each field's `edge` says whether it opens a group or a sub-group,
	for the dividers."""
	groups = []
	for col in columns:
		sq = col.get("sq_link")
		if not sq:
			continue
		if not groups or groups[-1]["sq"] != sq:
			groups.append({
				"display": col.get("supplier_group") or sq,
				"sq": sq,
				"fields": [],
				"subgroups": [],
				"qty_field": None,
				"rate_field": None,
				"amount_field": None,
				"narration_field": None,
			})
		group = groups[-1]

		kind = next(
			(name for suffix, name in _SUPPLIER_FIELD_KINDS if col["fieldname"].endswith(suffix)),
			"amount",
		)
		field = {
			"kind": kind,
			"field": col["fieldname"],
			"label": col.get("label") or "",
			"po": col.get("po_link"),
			"edge": "",
		}
		group["fields"].append(field)
		if kind in ("qty", "rate", "amount", "narration"):
			group[kind + "_field"] = col["fieldname"]

		label = col.get("sub_group") or ""
		if not group["subgroups"] or group["subgroups"][-1]["label"] != label:
			group["subgroups"].append({"label": label, "po": col.get("po_link"), "fields": []})
		group["subgroups"][-1]["fields"].append(field)

	for index, group in enumerate(groups):
		group["span"] = len(group["fields"])
		group["band"] = SUPPLIER_BAND_COLORS[index % len(SUPPLIER_BAND_COLORS)]
		for position, sub in enumerate(group["subgroups"]):
			sub["span"] = len(sub["fields"])
			sub["fields"][0]["edge"] = "grp-start" if position == 0 else "sub-start"
	return groups


@frappe.whitelist()
def export_xlsx(filters):
	"""Excel export laid out like the on-screen comparison, with the same three
	header rows merged like a hand-made sheet: the Supplier Quotation across its
	whole block; under it "Quoted" and each Purchase Order made from that quotation
	across their own columns; under those Qty / Rate / Amount. Quotation blocks
	carry the report's alternating tint and a divider on the left edge, and the
	quotation and PO headings link to their documents."""
	from io import BytesIO

	from frappe.utils.xlsxutils import make_xlsx
	from openpyxl import load_workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.properties import PageSetupProperties

	if isinstance(filters, str):
		filters = frappe._dict(json.loads(filters))

	columns, data = execute(filters)[:2]
	groups = _supplier_groups(columns)

	fixed_labels = ["SN", "Item Name", "MR Qty"]
	FIXED = len(fixed_labels)
	HEAD = 3  # header rows: quotation / Quoted or PO / Qty, Rate, Amount

	group_row = list(fixed_labels)  # fixed labels sit in row 1, merged down all three
	sub_row = [""] * FIXED
	label_row = [""] * FIXED
	for g in groups:
		group_row += [f"{g['display']}\n{g['sq']}"] + [""] * (g["span"] - 1)
		for sub in g["subgroups"]:
			sub_row += [sub["label"]] + [""] * (sub["span"] - 1)
			label_row += [f["label"] for f in sub["fields"]]

	rows = [group_row, sub_row, label_row]
	row_kinds = ["header"] * HEAD  # per sheet row, for the styling pass below
	for d in data:
		if not isinstance(d, dict) or not d:
			continue
		row = [d.get("sn"), d.get("item_name") or d.get("item_code"), d.get("qty")]
		for g in groups:
			row += [d.get(f["field"]) for f in g["fields"]]
		rows.append(row)
		if d.get("is_data_row"):
			row_kinds.append("data")
		elif d.get("is_term_row"):
			row_kinds.append("term")
		elif d.get("is_total_row") or d.get("is_invoice_row"):
			row_kinds.append("emphasis")
		else:
			row_kinds.append("summary")

	xlsx = make_xlsx(rows, "Supplier Quotation Comparison")

	# make_xlsx works on a write-only workbook, so restyle by reopening the file.
	wb = load_workbook(xlsx)
	ws = wb.active
	last_row = ws.max_row
	last_col = len(label_row)

	# Nepali digit grouping (1,98,750.00), matching what the report and the PDF show.
	MONEY = r"[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;#,##0.00"
	QTY = "#,##0.##"

	hair = Side(style="thin", color="B0B0B0")
	thick = Side(style="medium", color="6B7280")  # left edge of a quotation block
	sub_edge = Side(style="thin", color="6B7280")  # left edge of a PO sub-group

	# Which sheet column each block / sub-group starts at, and what kind each column is.
	col_kind = {}  # 1-based sheet column -> qty / rate / amount / narration / poqty / porate / poamt
	block_starts, sub_starts = set(), set()
	col = FIXED + 1
	for g in groups:
		block_starts.add(col)
		for sub in g["subgroups"]:
			sub_starts.add(col)
			for offset, f in enumerate(sub["fields"]):
				col_kind[col + offset] = f["kind"]
			col += sub["span"]

	def left_edge(c):
		return thick if c in block_starts else (sub_edge if c in sub_starts else hair)

	# --- the three header rows ------------------------------------------------
	for c in range(1, FIXED + 1):
		ws.merge_cells(start_row=1, start_column=c, end_row=HEAD, end_column=c)
		cell = ws.cell(row=1, column=c)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)

	col = FIXED + 1
	for g in groups:
		ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + g["span"] - 1)
		cell = ws.cell(row=1, column=col)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.hyperlink = frappe.utils.get_url(f"/app/supplier-quotation/{g['sq']}")
		for sub in g["subgroups"]:
			if sub["span"] > 1:
				ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + sub["span"] - 1)
			cell = ws.cell(row=2, column=col)
			cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
			if sub["po"]:
				cell.hyperlink = frappe.utils.get_url(f"/app/purchase-order/{sub['po']}")
				cell.font = Font(bold=True, color="0563C1")
			else:
				cell.font = Font(bold=True)
			col += sub["span"]

	for c in range(FIXED + 1, last_col + 1):
		head = ws.cell(row=HEAD, column=c)
		head.font = Font(bold=True)
		head.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
	ws.row_dimensions[1].height = 32
	ws.row_dimensions[2].height = 18
	ws.row_dimensions[3].height = 18

	# --- body ----------------------------------------------------------------
	for r in range(1, last_row + 1):
		kind = row_kinds[r - 1] if r - 1 < len(row_kinds) else "data"
		for c in range(1, last_col + 1):
			cell = ws.cell(row=r, column=c)
			cell.border = Border(left=left_edge(c), right=hair, top=hair, bottom=hair)

			if r <= HEAD:
				continue

			column_kind = col_kind.get(c)
			if kind == "emphasis":
				cell.font = Font(bold=True)

			if kind == "term":
				# Free text (Payment Terms, Delivery Period, ...) - let it wrap
				# inside the column rather than run under the next supplier.
				cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
			elif column_kind in ("rate", "amount", "porate", "poamt"):
				cell.number_format = MONEY
				cell.alignment = Alignment(horizontal="right", vertical="center")
			elif column_kind in ("qty", "poqty"):
				cell.number_format = QTY
				cell.alignment = Alignment(horizontal="right", vertical="center")
			elif column_kind == "narration":
				cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
			elif c == 1:
				cell.alignment = Alignment(horizontal="right", vertical="center")
			elif c == 3 and kind == "data":
				cell.number_format = QTY
				cell.alignment = Alignment(horizontal="right", vertical="center")
			else:
				cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

	# --- merges, as the PDF lays them out ------------------------------------
	# A summary / term label belongs across the three fixed columns, not squeezed
	# into MR Qty; a term's free text belongs across its whole quotation block
	# rather than under whichever single column happened to carry it.
	for r in range(HEAD + 1, last_row + 1):
		kind = row_kinds[r - 1] if r - 1 < len(row_kinds) else "data"
		if kind == "data":
			continue

		label = ws.cell(row=r, column=3).value
		ws.cell(row=r, column=1).value = label
		ws.cell(row=r, column=3).value = None
		ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=FIXED)
		ws.cell(row=r, column=1).alignment = Alignment(
			horizontal="left", vertical="center", wrap_text=True
		)
		if kind == "emphasis":
			ws.cell(row=r, column=1).font = Font(bold=True)

		if kind != "term":
			continue
		col = FIXED + 1
		for g in groups:
			# The value sits in the block's Amount column; merging keeps only the
			# top-left cell, so carry it to the block's first column first.
			amount_offset = next(
				(i for i, f in enumerate(g["fields"]) if f["kind"] == "amount"), 0
			)
			value = ws.cell(row=r, column=col + amount_offset).value
			for offset in range(g["span"]):
				ws.cell(row=r, column=col + offset).value = None
			ws.cell(row=r, column=col).value = value
			if g["span"] > 1:
				ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + g["span"] - 1)
			ws.cell(row=r, column=col).alignment = Alignment(
				horizontal="left", vertical="top", wrap_text=True
			)
			col += g["span"]

	# Quotation tints last, so they sit under every cell of the block.
	col = FIXED + 1
	for g in groups:
		fill = PatternFill("solid", fgColor=g["band"].lstrip("#").upper())
		for offset in range(g["span"]):
			for r in range(1, last_row + 1):
				ws.cell(row=r, column=col + offset).fill = fill
		col += g["span"]

	# --- column widths -------------------------------------------------------
	# Sized from the widest value actually in each column (header rows 1-2 are
	# merged and would skew it), so the sheet opens readable. Capped, because a
	# long narration or item name would otherwise push the money columns off
	# screen; those columns wrap instead.
	CAPS = {"narration": 34, "rate": 15, "amount": 16, "qty": 11,
			"porate": 15, "poamt": 16, "poqty": 11}
	for c in range(1, last_col + 1):
		longest = 0
		for r in range(HEAD, last_row + 1):
			value = ws.cell(row=r, column=c).value
			if value is None:
				continue
			if isinstance(value, (int, float)):
				# what it will look like once the number format is applied
				text = f"{value:,.2f}"
			else:
				text = max(str(value).split("\n"), key=len)
			longest = max(longest, len(text))

		kind = col_kind.get(c)
		if c == 1:
			width = 5
		elif c == 2:
			width = min(max(longest + 2, 16), 30)
		elif c == 3:
			width = 9
		else:
			width = min(max(longest + 2, 10), CAPS.get(kind, 16))
		ws.column_dimensions[get_column_letter(c)].width = width

	# Item / label column and all three header rows stay put while scrolling.
	ws.freeze_panes = f"D{HEAD + 1}"

	# Printing from Excel should need no setup either.
	ws.print_title_rows = f"1:{HEAD}"
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

	out = BytesIO()
	wb.save(out)

	frappe.response.filename = "supplier_quotation_comparison.xlsx"
	frappe.response.filecontent = out.getvalue()
	frappe.response.type = "binary"


@frappe.whitelist()
def download_pdf(filters, view=None):
	"""Print / PDF of the comparison in document form: company + title header,
	one column block per supplier (quotation number under the supplier's name)
	holding the same columns in the same order as the on-screen report, and the
	same summary rows below, under the same three header rows (quotation / Quoted
	and each Purchase Order / Qty, Rate, Amount).
	Both menu entries call this — Print opens it inline (view=1), PDF downloads."""
	from frappe.utils.pdf import get_pdf

	from avinashgroup_app.custom_code.printing.chrome_pdf import render as chrome_render

	if isinstance(filters, str):
		filters = frappe._dict(json.loads(filters))

	columns, data = execute(filters)[:2]
	groups = _supplier_groups(columns)

	purchase_orders = _as_list(filters.get("purchase_order"))
	report_title = "Supplier Quotation Comparison"
	if purchase_orders:
		report_title += " — {0}".format(", ".join(purchase_orders))

	subline_parts = []
	material_requests = _as_list(filters.get("material_request"))
	if purchase_orders:
		material_requests += get_material_requests_from_purchase_order(purchase_orders)
	material_requests = list(dict.fromkeys(mr for mr in material_requests if mr))
	if material_requests:
		subline_parts.append("Material Request: {0}".format(", ".join(material_requests)))
	if cint(filters.get("preferred_quotation")):
		subline_parts.append("Preferred quotations only")

	currency = frappe.get_cached_value("Company", filters.get("company"), "default_currency")

	def fmt(value):
		return fmt_money(flt(value), 2, currency) if value is not None else ""

	template_path = os.path.join(os.path.dirname(__file__), "supplier_quotation_comparison_pdf.html")
	with open(template_path) as f:
		template_content = f.read()

	html = frappe.render_template(
		template_content,
		{
			"company": filters.get("company"),
			"report_title": report_title,
			"subline_parts": subline_parts,
			"groups": groups,
			"data": data,
			"fmt": fmt,
			"fmt_qty": fmt_qty,
			# Wide enough for a formatted NPR amount at the table's font size -
			# "Rs 1,98,750.00" plus cell padding. Too narrow and the money cells,
			# which must not wrap, spill over the column border into their neighbour.
			"field_widths": {
				"qty": "40px",
				"rate": "72px",
				"amount": "88px",
				"narration": "95px",
				"poqty": "40px",
				"porate": "72px",
				"poamt": "88px",
			},
		},
	)

	# The three fixed columns plus about six supplier columns sit on portrait A4;
	# beyond that they get too narrow, so use landscape. Counting the real columns
	# (rather than the suppliers) keeps this right whatever each block holds.
	total_columns = 3 + sum(g["span"] for g in groups)
	orientation = "Portrait" if total_columns <= 9 else "Landscape"

	# Page geometry has to be stated twice because the two renderers read it from
	# different places: wkhtmltopdf takes the options dict below, Chrome honours
	# only an @page rule in the document itself.
	html = (
		"<style>@page { size: A4 %s; margin: 10mm 10mm 12mm 10mm; }</style>"
		% orientation.lower()
	) + html

	# This app renders print formats with headless Chrome (hooks.py: pdf_generator),
	# and a bench that has Chrome generally has no wkhtmltopdf at all - calling
	# get_pdf first would raise "No wkhtmltopdf executable found" before anything
	# else got a chance. So try Chrome, and fall back only when it isn't there.
	pdf_data = chrome_render(html=html, pdf_generator="chrome")
	if pdf_data is None:
		pdf_data = get_pdf(html, {
			"page-size": "A4",
			"orientation": orientation,
			"margin-top": "10mm",
			"margin-right": "10mm",
			"margin-bottom": "12mm",
			"margin-left": "10mm",
			"encoding": "UTF-8",
		})

	frappe.response.filename = "supplier_quotation_comparison.pdf"
	frappe.response.filecontent = pdf_data
	# view=1 (Print) -> open inline in the browser tab; otherwise download the file.
	frappe.response.type = "pdf" if cint(view) else "download"


# ============================================
# Utility Functions
# ============================================

@frappe.whitelist()
def set_default_supplier(item_code, supplier, company):
	"""Set default supplier for an item"""
	frappe.db.set_value(
		"Item Default",
		{"parent": item_code, "company": company},
		"default_supplier",
		supplier,
	)


@frappe.whitelist()
def get_material_requests_from_purchase_order(purchase_order):
	"""Material Request(s) one or more Purchase Orders were raised from, read off
	their item lines.

	The approval workflow runs on the Purchase Order; approvers open the comparison
	from there. A PO may draw on more than one Material Request, and the report can
	compare several POs at once, so this takes a single name or a list (plain or
	JSON-encoded) and returns the union; get_data() matches Supplier Quotation Items
	against all of them.
	"""
	purchase_orders = _as_list(purchase_order)
	if not purchase_orders:
		return []
	rows = frappe.db.sql(
		"""
		SELECT DISTINCT poi.material_request
		FROM `tabPurchase Order Item` poi
		WHERE poi.parent IN %(pos)s AND poi.material_request IS NOT NULL
		""",
		{"pos": tuple(purchase_orders)},
	)
	return [r[0] for r in rows]